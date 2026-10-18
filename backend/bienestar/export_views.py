"""US32: exportar los hábitos nutricionales y las estadísticas del usuario en Excel (.xlsx)."""

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from rest_framework.views import APIView

from .models import ComidaPlan, PlanNutricional, ProgresoDiario, RutinaEjercicio
from .nutrition_views import UsuarioSessionAuthentication

TIPO_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
SIN_DATOS = 'Sin datos registrados'
FORMATO_FECHA = 'DD/MM/YYYY'

# Colores de la app: oscuro para el título, turquesa para las secciones y celeste para los encabezados.
COLOR_TITULO = '2C3E50'
COLOR_SECCION = '7FFFD4'
COLOR_ENCABEZADO = 'D6EEF8'
BORDE = Border(*(Side(style='thin', color='B0BEC5'),) * 4)
ANCHOS = {'A': 24, 'B': 34, 'C': 18, 'D': 24, 'E': 16}


class HojaDeDatos:
    """Escribe la hoja fila por fila, con un estilo distinto para títulos, encabezados y datos."""

    def __init__(self, hoja):
        self.hoja = hoja
        self.fila = 1

    def _celda(self, columna, valor, formato=None):
        celda = self.hoja.cell(row=self.fila, column=columna, value=valor)
        if isinstance(valor, str) and valor.startswith('='):
            # Un texto escrito por el usuario nunca debe ejecutarse como fórmula.
            celda.data_type = 's'
        if formato:
            celda.number_format = formato
        return celda

    def titulo(self, texto):
        celda = self._celda(1, texto)
        celda.font = Font(bold=True, size=16, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=COLOR_TITULO)
        self.hoja.merge_cells(start_row=self.fila, start_column=1, end_row=self.fila, end_column=len(ANCHOS))
        self.hoja.row_dimensions[self.fila].height = 26
        self.fila += 1

    def seccion(self, texto):
        self.fila += 1  # fila vacía de separación
        celda = self._celda(1, texto)
        celda.font = Font(bold=True, size=12)
        celda.fill = PatternFill('solid', fgColor=COLOR_SECCION)
        self.hoja.merge_cells(start_row=self.fila, start_column=1, end_row=self.fila, end_column=len(ANCHOS))
        self.fila += 1

    def encabezados(self, *textos):
        for columna, texto in enumerate(textos, start=1):
            celda = self._celda(columna, texto)
            celda.font = Font(bold=True)
            celda.fill = PatternFill('solid', fgColor=COLOR_ENCABEZADO)
            celda.border = BORDE
        self.fila += 1

    def datos(self, *valores):
        """Cada valor es un dato simple o una tupla (dato, formato de número de Excel)."""
        for columna, valor in enumerate(valores, start=1):
            dato, formato = valor if isinstance(valor, tuple) else (valor, None)
            celda = self._celda(columna, dato, formato)
            celda.border = BORDE
            celda.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        self.fila += 1

    def etiqueta(self, texto, valor, formato=None):
        self._celda(1, texto).font = Font(bold=True)
        self._celda(2, valor, formato).alignment = Alignment(horizontal='left')
        self.fila += 1

    def nota(self, texto):
        self._celda(1, texto).font = Font(italic=True, color='607D8B')
        self.fila += 1


def escribir_plan_nutricional(hoja, usuario):
    hoja.seccion('Plan nutricional')
    plan = PlanNutricional.objects.filter(usuario=usuario).prefetch_related('comidas').first()
    if plan is None:
        hoja.nota(SIN_DATOS)
        return

    hoja.etiqueta('Nombre', plan.nombre)
    hoja.etiqueta('Objetivo', plan.objetivo or '-')
    if plan.notas:
        hoja.etiqueta('Notas', plan.notas)
    hoja.etiqueta('Última actualización', timezone.localtime(plan.actualizado).date(), FORMATO_FECHA)

    comidas = plan.comidas.all()
    if comidas:
        hoja.fila += 1
        hoja.encabezados('Comida', 'Descripción', 'Horario')
        for comida in comidas:
            hoja.datos(
                ComidaPlan.Tipo(comida.tipo).label,
                comida.descripcion,
                comida.horario.strftime('%H:%M') if comida.horario else '-',
            )


def escribir_habitos_diarios(hoja, usuario):
    hoja.seccion('Hábitos diarios')
    progresos = (
        ProgresoDiario.objects.filter(usuario=usuario)
        .select_related('habito')
        .order_by('-fecha', 'habito__nombre')
    )
    if not progresos:
        hoja.nota(SIN_DATOS)
        return

    hoja.encabezados('Fecha', 'Hábito', 'Registro')
    for progreso in progresos:
        if progreso.valor is not None:
            registro = (float(progreso.valor), 'General" horas"')
        else:
            registro = 'Cumplido' if progreso.completado else 'No cumplido'
        hoja.datos((progreso.fecha, FORMATO_FECHA), progreso.habito.nombre, registro)


def escribir_rutina_y_estadisticas(hoja, usuario):
    rutinas = list(
        RutinaEjercicio.objects.filter(usuario=usuario)
        .select_related('ejercicio')
        .order_by('-fecha_registro', 'ejercicio__nombre')
    )

    hoja.seccion('Rutina de ejercicios')
    if rutinas:
        hoja.encabezados('Fecha', 'Ejercicio', 'Tipo', 'Meta', 'Estado')
        for rutina in rutinas:
            hoja.datos(
                (rutina.fecha_registro, FORMATO_FECHA),
                rutina.ejercicio.nombre,
                rutina.ejercicio.get_tipo_display(),
                (rutina.meta_cantidad, '0" repeticiones/km"'),
                'Completado' if rutina.completado else 'Pendiente',
            )
    else:
        hoja.nota(SIN_DATOS)

    hoja.seccion('Estadísticas')
    if not rutinas:
        hoja.nota(SIN_DATOS)
        return

    total = len(rutinas)
    completadas = sum(rutina.completado for rutina in rutinas)
    hoja.encabezados('Indicador', 'Valor')
    hoja.datos('Ejercicios registrados', total)
    hoja.datos('Ejercicios completados', completadas)
    hoja.datos('Cumplimiento', (completadas / total, '0%'))


def generar_excel(usuario, hoy):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Mis datos'
    for columna, ancho in ANCHOS.items():
        hoja.column_dimensions[columna].width = ancho
    # Al imprimir o guardar como PDF, todas las columnas entran en el ancho de la hoja.
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0

    escritor = HojaDeDatos(hoja)
    escritor.titulo('Bienestar Plus - Mis datos')
    escritor.etiqueta('Usuario', usuario.nombre)
    escritor.etiqueta('Fecha de exportación', hoy, FORMATO_FECHA)
    escribir_plan_nutricional(escritor, usuario)
    escribir_habitos_diarios(escritor, usuario)
    escribir_rutina_y_estadisticas(escritor, usuario)

    archivo = BytesIO()
    libro.save(archivo)
    return archivo.getvalue()


class ExportarDatosView(APIView):
    """
    GET /api/exportar-datos/
    Descarga un Excel con los datos del usuario en sesión. Nunca recibe un id por
    parámetro: solo puede exportar sus propios datos.
    """

    authentication_classes = [UsuarioSessionAuthentication]

    def get(self, request):
        hoy = timezone.localdate()
        respuesta = HttpResponse(generar_excel(request.user, hoy), content_type=TIPO_XLSX)
        respuesta['Content-Disposition'] = f'attachment; filename="bienestar_mis_datos_{hoy:%Y-%m-%d}.xlsx"'
        respuesta['Cache-Control'] = 'no-store'
        return respuesta