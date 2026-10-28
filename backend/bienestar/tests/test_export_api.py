"""US32: exportar hábitos nutricionales y estadísticas en Excel (.xlsx)."""

from datetime import datetime, time
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.hashers import make_password
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from bienestar.models import (
    ComidaPlan, Ejercicio, Habito, PlanNutricional, ProgresoDiario,
    Roles, RutinaEjercicio, Tipos, Usuario,
)

SECCIONES = ('Plan nutricional', 'Hábitos diarios', 'Rutina de ejercicios', 'Estadísticas')


class ExportarDatosApiTests(TestCase):
    def setUp(self):
        self.url = reverse('exportar-datos')
        self.client = APIClient()
        self.usuario = Usuario.objects.create(
            nombre='Ana QA', email='ana.export@example.test',
            password=make_password('clave-ana'), rol=Roles.ESTANDAR,
        )
        self.otro = Usuario.objects.create(
            nombre='Otro QA', email='otro.export@example.test',
            password=make_password('clave-otro'), rol=Roles.ESTANDAR,
        )

    def login(self, usuario, clave):
        response = self.client.post('/api/login/', {'email': usuario.email, 'password': clave}, format='json')
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def hoja(self, response):
        return load_workbook(BytesIO(response.content))['Mis datos']

    def filas(self, hoja):
        """Valores de cada fila, sin las celdas vacías del final."""
        filas = []
        for fila in hoja.iter_rows(values_only=True):
            # Excel guarda las fechas como fecha y hora: se comparan solo por la fecha.
            valores = [v.date() if isinstance(v, datetime) else v for v in fila]
            while valores and valores[-1] is None:
                valores.pop()
            filas.append(valores)
        return filas

    def seccion(self, filas, titulo):
        """Filas de una sección: desde su título hasta la próxima sección."""
        inicio = filas.index([titulo]) + 1
        fin = next((i for i in range(inicio, len(filas)) if filas[i][:1] and filas[i][0] in SECCIONES), len(filas))
        return [fila for fila in filas[inicio:fin] if fila]

    def cargar_datos(self, usuario, plan='Plan de Ana'):
        plan_obj = PlanNutricional.objects.create(usuario=usuario, nombre=plan, objetivo='Comer más verduras')
        ComidaPlan.objects.create(plan=plan_obj, tipo='desayuno', descripcion='Avena con fruta', horario=time(8, 30))
        sueno, _ = Habito.objects.get_or_create(nombre='Horas de sueño')
        agua, _ = Habito.objects.get_or_create(nombre='Hidratación')
        hoy = timezone.localdate()
        ProgresoDiario.objects.create(usuario=usuario, habito=sueno, fecha=hoy, valor=Decimal('7.5'), completado=True)
        ProgresoDiario.objects.create(usuario=usuario, habito=agua, fecha=hoy, completado=False)
        ejercicio = Ejercicio.objects.create(nombre=f'Sentadillas {usuario.pk}', tipo=Tipos.FUERZA)
        RutinaEjercicio.objects.create(usuario=usuario, ejercicio=ejercicio, meta_cantidad=20, completado=True)

    def test_requiere_sesion(self):
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_descarga_excel_con_nombre_de_archivo_y_cabecera(self):
        self.login(self.usuario, 'clave-ana')

        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertRegex(response['Content-Disposition'], r'attachment; filename="bienestar_mis_datos_[\d-]+\.xlsx"')
        filas = self.filas(self.hoja(response))
        self.assertEqual(filas[0], ['Bienestar Plus - Mis datos'])
        self.assertIn(['Usuario', 'Ana QA'], filas)
        self.assertIn(['Fecha de exportación', timezone.localdate()], filas)

    def test_cada_seccion_tiene_sus_columnas_y_datos(self):
        self.cargar_datos(self.usuario)
        self.login(self.usuario, 'clave-ana')

        filas = self.filas(self.hoja(self.client.get(self.url)))

        hoy = timezone.localdate()
        plan = self.seccion(filas, 'Plan nutricional')
        self.assertIn(['Nombre', 'Plan de Ana'], plan)
        self.assertIn(['Objetivo', 'Comer más verduras'], plan)
        self.assertIn(['Comida', 'Descripción', 'Horario'], plan)
        self.assertIn(['Desayuno', 'Avena con fruta', '08:30'], plan)

        habitos = self.seccion(filas, 'Hábitos diarios')
        self.assertEqual(habitos[0], ['Fecha', 'Hábito', 'Registro'])
        self.assertIn([hoy, 'Horas de sueño', 7.5], habitos)
        self.assertIn([hoy, 'Hidratación', 'No cumplido'], habitos)

        rutina = self.seccion(filas, 'Rutina de ejercicios')
        self.assertEqual(rutina[0], ['Fecha', 'Ejercicio', 'Tipo', 'Meta', 'Estado'])
        self.assertIn([hoy, f'Sentadillas {self.usuario.pk}', 'Fuerza', 20, 'Completado'], rutina)

        estadisticas = self.seccion(filas, 'Estadísticas')
        self.assertEqual(estadisticas[0], ['Indicador', 'Valor'])
        self.assertIn(['Ejercicios registrados', 1], estadisticas)
        self.assertIn(['Ejercicios completados', 1], estadisticas)
        self.assertIn(['Cumplimiento', 1], estadisticas)

    def test_formato_legible_fechas_horas_y_porcentaje(self):
        self.cargar_datos(self.usuario)
        self.login(self.usuario, 'clave-ana')

        hoja = self.hoja(self.client.get(self.url))

        celdas = {celda.value: celda for fila in hoja.iter_rows() for celda in fila if celda.value is not None}
        self.assertTrue(celdas['Bienestar Plus - Mis datos'].font.bold)
        self.assertTrue(celdas['Fecha'].font.bold)  # encabezado de tabla
        self.assertGreaterEqual(hoja.column_dimensions['B'].width, 30)
        fila_sueno = next(fila for fila in hoja.iter_rows() if fila[1].value == 'Horas de sueño')
        self.assertEqual(fila_sueno[0].number_format, 'DD/MM/YYYY')
        self.assertEqual(fila_sueno[2].number_format, 'General" horas"')
        fila_cumplimiento = next(fila for fila in hoja.iter_rows() if fila[0].value == 'Cumplimiento')
        self.assertEqual(fila_cumplimiento[1].number_format, '0%')

    def test_solo_exporta_los_datos_del_usuario_en_sesion(self):
        self.cargar_datos(self.usuario, plan='Plan de Ana')
        self.cargar_datos(self.otro, plan='Plan del otro')
        self.login(self.usuario, 'clave-ana')

        valores = {v for fila in self.filas(self.hoja(self.client.get(self.url))) for v in fila}

        self.assertIn('Plan de Ana', valores)
        self.assertNotIn('Plan del otro', valores)
        self.assertNotIn(f'Sentadillas {self.otro.pk}', valores)

    def test_caso_sin_datos(self):
        self.login(self.usuario, 'clave-ana')

        filas = self.filas(self.hoja(self.client.get(self.url)))

        for titulo in SECCIONES:
            self.assertEqual(self.seccion(filas, titulo)[0], ['Sin datos registrados'])

    def test_texto_del_usuario_no_se_ejecuta_como_formula(self):
        PlanNutricional.objects.create(usuario=self.usuario, nombre='=HYPERLINK("http://x")', objetivo='ok')
        self.login(self.usuario, 'clave-ana')

        hoja = self.hoja(self.client.get(self.url))

        celda = next(c for fila in hoja.iter_rows() for c in fila if c.value == '=HYPERLINK("http://x")')
        self.assertEqual(celda.data_type, 's')